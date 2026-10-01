from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from typing import Optional
from pydantic import BaseModel
from sqlalchemy import select, func, or_
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
from app.database import get_db
from app.models import Guest, Booking, Room, RoomType, Payment
from app.dependencies import get_current_user
from app.serialize import row_to_dict
from app.services.phone import normalize_phone
from enum import Enum
import io
import math


class IdType(str, Enum):
    AADHAAR = "AADHAAR"
    PAN = "PAN"
    PASSPORT = "PASSPORT"
    DRIVING_LICENSE = "DRIVING_LICENSE"
    VOTER_ID = "VOTER_ID"


class CreateGuestRequest(BaseModel):
    firstName: str
    lastName: str
    phone: str
    email: Optional[str] = None
    address: Optional[str] = None
    city: Optional[str] = None
    state: Optional[str] = None
    idType: Optional[str] = None
    idNumber: Optional[str] = None
    dateOfBirth: Optional[str] = None
    anniversary: Optional[str] = None
    notes: Optional[str] = None
    source: Optional[str] = None
    whatsappConsent: Optional[bool] = True


class UpdateGuestRequest(BaseModel):
    firstName: Optional[str] = None
    lastName: Optional[str] = None
    phone: Optional[str] = None
    email: Optional[str] = None
    address: Optional[str] = None
    city: Optional[str] = None
    state: Optional[str] = None
    idType: Optional[IdType] = None
    idNumber: Optional[str] = None
    idDocumentUrl: Optional[str] = None
    dateOfBirth: Optional[str] = None
    anniversary: Optional[str] = None
    notes: Optional[str] = None
    source: Optional[str] = None
    whatsappConsent: Optional[bool] = None
    dnd: Optional[bool] = None


router = APIRouter(prefix="/api/guests", tags=["Guests"])


@router.post("")
async def create_guest(body: CreateGuestRequest, user=Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    from datetime import datetime as dt
    phone = normalize_phone(body.phone)
    if not phone or not 8 <= len(phone) <= 15:
        raise HTTPException(400, "Enter a phone number with country code")
    existing = await db.execute(select(Guest).where(Guest.phone == phone, Guest.deletedAt.is_(None)))
    if existing.scalar_one_or_none():
        raise HTTPException(409, "Guest with this phone already exists")
    guest = Guest(firstName=body.firstName, lastName=body.lastName, phone=phone, source=body.source or "WALK_IN")
    if body.email:
        guest.email = body.email
    if body.address:
        guest.address = body.address
    if body.city:
        guest.city = body.city
    if body.state:
        guest.state = body.state
    if body.idType:
        guest.idType = body.idType
    if body.idNumber:
        guest.idNumber = body.idNumber
    if body.notes:
        guest.notes = body.notes
    if body.whatsappConsent is not None:
        guest.whatsappConsent = body.whatsappConsent
    if body.dateOfBirth:
        guest.dateOfBirth = dt.fromisoformat(body.dateOfBirth)
    if body.anniversary:
        guest.anniversary = dt.fromisoformat(body.anniversary)
    db.add(guest)
    await db.commit()
    await db.refresh(guest)
    return row_to_dict(guest)


@router.get("")
async def get_guests(
    page: int = Query(1, ge=1),
    limit: int = Query(20, ge=1, le=100),
    user=Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    skip = (page - 1) * limit
    base = select(Guest).where(Guest.deletedAt.is_(None))
    total_q = await db.execute(select(func.count()).select_from(base.subquery()))
    total = total_q.scalar()
    result = await db.execute(base.order_by(Guest.createdAt.desc()).offset(skip).limit(limit))
    data = [row_to_dict(g) for g in result.scalars().all()]
    return {
        "data": data,
        "total": total,
        "page": page,
        "limit": limit,
        "totalPages": math.ceil(total / limit) if total > 0 else 0,
        "meta": {"total": total, "page": page, "limit": limit, "totalPages": math.ceil(total / limit) if total > 0 else 0},
    }


@router.post("/import")
async def import_guests(
    file: UploadFile = File(...),
    commit: bool = Query(False),
    user=Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    raw = await file.read()
    name = (file.filename or "").lower()
    rows = []
    if name.endswith(".csv"):
        import csv
        text = raw.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
    else:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for excel_row in ws.iter_rows(min_row=2, values_only=True):
            rows.append({headers[i]: excel_row[i] if i < len(excel_row) else None for i in range(len(headers))})

    def _cell(row, *keys):
        lower = {str(k).strip().lower(): v for k, v in row.items()}
        for k in keys:
            if k in lower and lower[k] not in (None, ""):
                return str(lower[k]).strip()
        return ""

    created = updated = skipped = 0
    errors = []
    preview = []
    for i, row in enumerate(rows, start=2):
        phone = normalize_phone(_cell(row, "phone", "mobile", "whatsapp"))
        first = _cell(row, "firstname", "first_name", "first name", "name")
        last = _cell(row, "lastname", "last_name", "last name")
        if first and not last and " " in first:
            first, last = first.split(" ", 1)
        if not phone:
            errors.append({"row": i, "error": "missing phone"})
            skipped += 1
            continue
        if not first:
            first = "Guest"
        email = _cell(row, "email")
        city = _cell(row, "city")
        notes = _cell(row, "notes", "note")
        existing = await db.execute(select(Guest).where(Guest.phone == phone, Guest.deletedAt.is_(None)))
        guest = existing.scalar_one_or_none()
        action = "update" if guest else "create"
        preview.append({"row": i, "action": action, "phone": phone, "firstName": first, "lastName": last})
        if not commit:
            if guest:
                updated += 1
            else:
                created += 1
            continue
        if guest:
            guest.firstName = first or guest.firstName
            guest.lastName = last if last != "" else guest.lastName
            if email:
                guest.email = email
            if city:
                guest.city = city
            if notes:
                guest.notes = ((guest.notes or "") + "\n" + notes).strip()
            updated += 1
        else:
            db.add(Guest(
                firstName=first,
                lastName=last,
                phone=phone,
                email=email or None,
                city=city or None,
                notes=notes or None,
                source="IMPORT",
                whatsappConsent=True,
            ))
            created += 1
    if commit:
        await db.commit()
    return {
        "commit": commit,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
        "preview": preview[:200],
        "totalRows": len(rows),
    }


@router.get("/search")
async def search_guests(
    query: str = Query(None),
    q: str = Query(None),
    user=Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    query = query or q
    if not query:
        raise HTTPException(400, "query required")
    result = await db.execute(
        select(Guest).where(
            Guest.deletedAt.is_(None),
            or_(
                Guest.firstName.ilike(f"%{query}%"),
                Guest.lastName.ilike(f"%{query}%"),
                Guest.phone.contains(query),
                Guest.email.ilike(f"%{query}%"),
                Guest.idNumber.contains(query),
            ),
        ).order_by(Guest.createdAt.desc()).limit(20)
    )
    return [row_to_dict(g) for g in result.scalars().all()]


@router.get("/{id}")
async def get_guest(id: str, user=Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    guest = await db.get(Guest, id)
    if not guest or guest.deletedAt:
        raise HTTPException(404, "Guest not found")
    return row_to_dict(guest)


@router.get("/{id}/history")
async def get_guest_history(id: str, user=Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    guest = await db.get(Guest, id)
    if not guest:
        raise HTTPException(404, "Guest not found")
    result = await db.execute(
        select(Booking)
        .options(selectinload(Booking.room).selectinload(Room.roomType), selectinload(Booking.payments))
        .where(Booking.guestId == id)
        .order_by(Booking.checkIn.desc())
    )
    return [row_to_dict(b, rels={"room", "payments", "guest"}) for b in result.scalars().all()]


@router.patch("/{id}")
async def update_guest(id: str, body: UpdateGuestRequest, user=Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    guest = await db.get(Guest, id)
    if not guest or guest.deletedAt:
        raise HTTPException(404, "Guest not found")
    data = body.model_dump(exclude_none=True)
    if "phone" in data:
        data["phone"] = normalize_phone(data["phone"])
        if not data["phone"] or not 8 <= len(data["phone"]) <= 15:
            raise HTTPException(400, "Enter a phone number with country code")
    if data.get("email") == "":
        data["email"] = None
    if "dateOfBirth" in data:
        from datetime import datetime as dt
        data["dateOfBirth"] = dt.fromisoformat(data["dateOfBirth"])
    if "anniversary" in data:
        from datetime import datetime as dt
        data["anniversary"] = dt.fromisoformat(data["anniversary"])
    for k, v in data.items():
        setattr(guest, k, v)
    await db.commit()
    await db.refresh(guest)
    return row_to_dict(guest)
